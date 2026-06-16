"""
Belge Ayrıştırıcı (Document Parser)
.docx dosyalarından yapısal bilgi çıkarır:
- Paragrafları bölümlere ayırır (başlık, sayı/tarih, konu, metin, imza vb.)
- Font, boyut, hizalama bilgilerini toplar
- Marj bilgilerini okur
"""
from __future__ import annotations
import re
import logging
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.models.finding import (
    ParsedDocument, ParsedParagraph, DocumentSection
)
from app.rules.hierarchy import FACULTIES, INSTITUTES, DEPARTMENTS
from app.services.turkish_text import tr_upper, tr_lower, is_closing_line

logger = logging.getLogger(__name__)


# ── Bölüm tespit kalıpları ──────────────────────────────────────────────
_RE_TC = re.compile(r"^\s*T\s*\.?\s*C\s*\.?\s*$", re.IGNORECASE)
_RE_UNIVERSITY = re.compile(
    r"GEBZE\s+TEKNİK\s+ÜNİVERSİTESİ\s+REKTÖRLÜĞÜ", re.IGNORECASE
)
_RE_SAYI = re.compile(r"^\s*Sayı\s*:", re.IGNORECASE)
_RE_TARIH = re.compile(r"\d{2}[./]\d{2}[./]\d{4}")
_RE_KONU = re.compile(r"^\s*Konu\s*:", re.IGNORECASE)
_RE_ILGI = re.compile(r"^\s*İlgi\s*:", re.IGNORECASE)
_RE_ILGI_ITEM = re.compile(r"^\s*[a-zçğıöşü]\)", re.IGNORECASE)
_RE_EK = re.compile(r"^\s*EK\s*:", re.IGNORECASE)
_RE_EK_ITEM = re.compile(r"^\s*EK[\s-]*\d", re.IGNORECASE)
_RE_DAGITIM = re.compile(r"^\s*DAĞITIM", re.IGNORECASE)
# Kapanış ifadesi tespiti app.services.turkish_text.is_closing_line() ile yapılır
# (çekim-toleranslı: "arz ederim", "arz edilmektedir", "arz olunur" vb.).
_RE_REKTOR_A = re.compile(r"Rektör\s+a\.", re.IGNORECASE)

_KNOWN_UNITS = [*FACULTIES, *INSTITUTES, *DEPARTMENTS]
# Türkçe-duyarlı büyük harf: "Bilgi İşlem".upper() yerine tr_upper kullanılır
# (i→İ, ı→I) ki karşılaştırma her iki tarafta da tutarlı olsun.
_KNOWN_UNITS_UPPER = [tr_upper(name) for name in _KNOWN_UNITS]


def _matches_known_unit(text: str) -> bool:
    upper_text = tr_upper(text)
    return any(unit in upper_text for unit in _KNOWN_UNITS_UPPER)


# Resmî yazıda muhatap (alıcı) satırı tipik olarak bir makam/birim ekiyle ve
# virgülle biter ("...Makamına,", "...Dekanlığına,") veya "Sayın ..." ile başlar.
_MUHATAP_SUFFIXES = (
    "makamına", "makamına,",
    "başkanlığına", "başkanlıklarına",
    "dekanlığına", "dekanlıklarına",
    "müdürlüğüne", "müdürlüklerine",
    "rektörlüğüne", "rektörlüğüne,",
    "bakanlığına", "valiliğine",
    "daire başkanlığına", "genel müdürlüğüne",
    "birimlere", "birimine", "birimlerine",
)


def _looks_like_muhatap(text: str) -> bool:
    """Bir satırın muhatap (alıcı) satırı olup olmadığını skorlu olarak değerlendirir.

    Tek bir konuma güvenmek yerine biçim + anahtar kelime ipuçlarını birleştirir:
      • aşırı uzun değil (≤ 200 karakter),
      • "Sayın ..." ile başlıyor, VEYA
      • bir alıcı ekiyle (Makamına/Dekanlığına/Başkanlığına...) ve genellikle
        virgülle bitiyor.
    """
    if not text or len(text) > 200:
        return False
    t = tr_lower(text.strip())
    if t.startswith("sayın "):
        return True
    stripped = t.rstrip(" ,.")
    return stripped.endswith(_MUHATAP_SUFFIXES)


def _get_doc_default_font(doc: Document) -> tuple[str | None, float | None]:
    """
    Belgenin varsayılan font adı ve boyutunu döndürür.
    Tema referanslarını (minorHAnsi, majorHAnsi) çözümler.
    Gerçek Word belgelerinde fontlar çoğunlukla stil/tema'dan miras alınır.
    """
    try:
        from docx.oxml.ns import qn
        from lxml import etree

        font_name: str | None = None
        size_pt: float | None = None

        # Tema font çözümlemesi için önce tema dosyasını oku
        minor_font: str | None = None
        major_font: str | None = None
        for rel in doc.part.rels.values():
            if "theme" in rel.reltype:
                ns = "http://schemas.openxmlformats.org/drawingml/2006/main"
                theme_xml = etree.fromstring(rel.target_part.blob)
                fontScheme = theme_xml.find(f".//{{{ns}}}fontScheme")
                if fontScheme is not None:
                    minor = fontScheme.find(f"{{{ns}}}minorFont")
                    major = fontScheme.find(f"{{{ns}}}majorFont")
                    if minor is not None:
                        lat = minor.find(f"{{{ns}}}latin")
                        if lat is not None:
                            minor_font = lat.get("typeface")
                    if major is not None:
                        lat = major.find(f"{{{ns}}}latin")
                        if lat is not None:
                            major_font = lat.get("typeface")
                break

        # docDefaults'tan varsayılan boyut ve font oku
        styles_el = doc.part.styles._element
        docDefaults = styles_el.find(qn("w:docDefaults"))
        if docDefaults is not None:
            rPrDefault = docDefaults.find(".//" + qn("w:rPrDefault"))
            if rPrDefault is not None:
                sz = rPrDefault.find(".//" + qn("w:sz"))
                if sz is not None:
                    val = sz.get(qn("w:val"))
                    if val:
                        size_pt = int(val) / 2  # yarım nokta → nokta

                rFonts = rPrDefault.find(".//" + qn("w:rFonts"))
                if rFonts is not None:
                    ascii_explicit = rFonts.get(qn("w:ascii"))
                    ascii_theme = rFonts.get(qn("w:asciiTheme")) or ""
                    if ascii_explicit:
                        font_name = ascii_explicit
                    elif "minor" in ascii_theme.lower() and minor_font:
                        font_name = minor_font
                    elif "major" in ascii_theme.lower() and major_font:
                        font_name = major_font

        return font_name, size_pt
    except Exception as exc:
        logger.debug("Varsayılan font okunamadı: %s", exc)
        return None, None


def _get_paragraph_font(
    para,
    fallback_font: str | None = None,
    fallback_size: float | None = None,
) -> tuple[str | None, float | None, bool | None]:
    """
    Paragraftaki baskın font adı, boyutu ve kalınlık durumunu döndürür.
    Run seviyesinde font yoksa (miras/tema durumu) belge varsayılanını kullanır.
    """
    fonts, sizes, bolds = [], [], []
    for run in para.runs:
        name = run.font.name
        size = run.font.size

        # Run'da explicit font yoksa karakter stiline bak
        if name is None and run.style and run.style.font.name:
            name = run.style.font.name
        if size is None and run.style and run.style.font.size:
            size = run.style.font.size

        if name:
            fonts.append(name)
        if size:
            sizes.append(size.pt)
        if run.font.bold is not None:
            bolds.append(run.font.bold)

    # Run'lardan font bulunamazsa paragraf stiline, oradan belge varsayılanına bak
    if not fonts:
        s = para.style
        while s is not None:
            if s.font.name:
                fonts.append(s.font.name)
                break
            s = s.base_style
        if not fonts and fallback_font:
            fonts.append(fallback_font)

    if not sizes:
        s = para.style
        while s is not None:
            if s.font.size:
                sizes.append(s.font.size.pt)
                break
            s = s.base_style
        if not sizes and fallback_size:
            sizes.append(fallback_size)

    font_name = max(set(fonts), key=fonts.count) if fonts else None
    font_size = max(set(sizes), key=sizes.count) if sizes else None
    is_bold = any(bolds) if bolds else None
    return font_name, font_size, is_bold


def _get_alignment_str(alignment) -> str | None:
    if alignment is None:
        return None
    mapping = {
        WD_ALIGN_PARAGRAPH.LEFT: "left",
        WD_ALIGN_PARAGRAPH.CENTER: "center",
        WD_ALIGN_PARAGRAPH.RIGHT: "right",
        WD_ALIGN_PARAGRAPH.JUSTIFY: "justify",
    }
    return mapping.get(alignment, str(alignment))


def _get_margins(doc: Document) -> dict | None:
    """Belgenin sayfa marjlarını cm cinsinden döndürür."""
    try:
        section = doc.sections[0]
        def emu_to_cm(val):
            if val is None:
                return None
            return round(val / 360000, 2)  # EMU → cm
        return {
            "top_cm": emu_to_cm(section.top_margin),
            "bottom_cm": emu_to_cm(section.bottom_margin),
            "left_cm": emu_to_cm(section.left_margin),
            "right_cm": emu_to_cm(section.right_margin),
        }
    except Exception:
        return None


def parse_docx(file_path: str | Path, original_filename: str | None = None) -> ParsedDocument:
    """
    .docx dosyasını ayrıştırır ve yapılandırılmış ParsedDocument döndürür.
    """
    doc = Document(str(file_path))
    filename = original_filename or Path(file_path).name

    parsed = ParsedDocument(filename=filename)
    parsed.page_margins = _get_margins(doc)

    # Belge varsayılan fontunu çözümle (miras/tema fontları için)
    doc_default_font, doc_default_size = _get_doc_default_font(doc)

    all_fonts = set()
    all_sizes = set()
    paragraphs: list[ParsedParagraph] = []

    # ── İlk geçiş: paragrafları topla ──
    for i, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        if not text:
            continue

        font_name, font_size, is_bold = _get_paragraph_font(
            para,
            fallback_font=doc_default_font,
            fallback_size=doc_default_size,
        )
        alignment = _get_alignment_str(para.alignment)
        is_centered = alignment == "center"
        
        if font_name:
            all_fonts.add(font_name)
        if font_size:
            all_sizes.add(font_size)
        
        pp = ParsedParagraph(
            index=i,
            text=text,
            font_name=font_name,
            font_size=font_size,
            is_bold=is_bold,
            is_centered=is_centered,
            alignment=alignment,
        )
        paragraphs.append(pp)
    
    parsed.fonts_used = sorted(all_fonts)
    parsed.font_sizes_used = sorted(all_sizes)
    
    # ── İkinci geçiş: bölüm tespiti ──
    current_section = DocumentSection.HEADER
    header_done = False
    sayi_tarih_done = False
    konu_done = False
    ilgi_active = False
    metin_started = False
    kapanis_done = False
    dagitim_active = False
    
    for pp in paragraphs:
        text = pp.text
        
        # T.C. başlığı
        if _RE_TC.match(text):
            pp.section = DocumentSection.HEADER
            parsed.has_tc_header = True
            continue
        
        # Üniversite adı
        if _RE_UNIVERSITY.search(text):
            pp.section = DocumentSection.HEADER
            parsed.has_university_name = True
            parsed.header_text = text
            continue
        
        # Birim adı (üniversite adından sonraki ortalanmış satır)
        if not header_done and pp.is_centered and not _RE_SAYI.match(text):
            if parsed.has_university_name and not parsed.unit_name:
                if not _RE_TC.match(text) and not _RE_UNIVERSITY.search(text):
                    pp.section = DocumentSection.HEADER
                    if _matches_known_unit(text):
                        parsed.unit_name = text
                    continue
        
        # Sayı satırı
        if _RE_SAYI.match(text):
            pp.section = DocumentSection.SAYI_TARIH
            parsed.sayi = text
            header_done = True
            sayi_tarih_done = True
            continue
        
        # Tarih (sayı satırında olmayıp ayrı satırda olabilir)
        if not sayi_tarih_done and _RE_TARIH.search(text) and len(text) < 30:
            pp.section = DocumentSection.SAYI_TARIH
            parsed.tarih = text
            continue
        
        # Konu
        if _RE_KONU.match(text):
            pp.section = DocumentSection.KONU
            parsed.konu = text
            konu_done = True
            header_done = True
            continue
        
        # Muhatap (alıcı) — skorlu tespit. Konu sonrası, metin başlamadan; İlgi
        # bloğundan SONRA gelse bile yakalanır (eski "ilk satır" sezgisi İlgi'den
        # sonra muhatabı kaçırıyordu). _looks_like_muhatap biçim + anahtar kelime
        # ipuçlarını birleştirir, böylece gövde cümleleri yanlışlıkla muhatap olmaz.
        if konu_done and not metin_started and parsed.muhatap is None:
            if _looks_like_muhatap(text):
                pp.section = DocumentSection.MUHATAP
                parsed.muhatap = text
                ilgi_active = False
                continue
        
        # İlgi
        if _RE_ILGI.match(text):
            pp.section = DocumentSection.ILGI
            parsed.ilgi_list.append(text)
            ilgi_active = True
            continue
        
        if ilgi_active and _RE_ILGI_ITEM.match(text):
            pp.section = DocumentSection.ILGI
            parsed.ilgi_list.append(text)
            continue
        else:
            ilgi_active = False
        
        # Kapanış ifadesi — kısa paragraflar kapanış olarak kabul edilir (< 70 karakter)
        # Uzun paragrafların sonu "arz ederim" ile bitebilir ama bunlar metin gövdesidir.
        # is_closing_line() çekim-toleranslıdır ("arz ederim" / "arz edilmektedir" /
        # "arz olunur" ...) ve eşleşmenin paragraf sonunda olmasını arar; böylece
        # "arz edilen konular" gibi gövde içi çekimler kapanış sayılmaz.
        if is_closing_line(text) and not kapanis_done and len(text) < 70:
            pp.section = DocumentSection.KAPANIS
            parsed.kapanis_phrase = text
            kapanis_done = True
            metin_started = False
            continue
        
        # Ek
        if _RE_EK.match(text) or _RE_EK_ITEM.match(text):
            pp.section = DocumentSection.EK
            parsed.ek_list.append(text)
            continue
        
        # Dağıtım başlığı ve alt satırları
        if _RE_DAGITIM.match(text):
            pp.section = DocumentSection.DAGITIM
            parsed.dagitim_list.append(text)
            dagitim_active = True
            continue

        if dagitim_active and not _RE_EK.match(text) and not _RE_EK_ITEM.match(text):
            # Dağıtım alt satırı (birim adları, Gereği/Bilgi başlıkları)
            pp.section = DocumentSection.DAGITIM
            parsed.dagitim_list.append(text)
            continue
        else:
            dagitim_active = False

        # İmza bloğu (kapanıştan sonra, ek'ten önce)
        if kapanis_done and not _RE_EK.match(text) and not _RE_DAGITIM.match(text):
            if len(text) < 150:
                pp.section = DocumentSection.IMZA
                if parsed.imza_block is None:
                    parsed.imza_block = text
                else:
                    parsed.imza_block += "\n" + text
                continue
        
        # Metin gövdesi (yukarıdaki hiçbir kalıba uymuyorsa)
        if header_done or konu_done:
            pp.section = DocumentSection.METIN
            parsed.metin_paragraphs.append(pp.index)
            metin_started = True
        else:
            # Hiçbir bölüme atanamadı — sessizce yanlış atama yapmak yerine paragrafı
            # UNKNOWN'da bırak ve logla (tanı/iyileştirme için izlenebilir).
            logger.debug("Parser: sınıflandırılamayan paragraf → unknown: %r", text[:60])
    
    # Tarih'i sayı satırından çıkar (eğer ayrı yoksa)
    if parsed.tarih is None and parsed.sayi:
        tarih_match = _RE_TARIH.search(parsed.sayi)
        if tarih_match:
            parsed.tarih = tarih_match.group(0)
    
    parsed.paragraphs = paragraphs
    return parsed
