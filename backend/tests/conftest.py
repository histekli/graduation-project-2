"""
Test fikstürleri — programatik olarak test .docx belgeleri oluşturur.
"""
from __future__ import annotations
import pytest
from pathlib import Path

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH


# ── Doküman oluşturucu yardımcıları ──────────────────────────────────────────

def _set_margins(doc: Document, cm: float = 1.5) -> None:
    for section in doc.sections:
        section.top_margin = Cm(cm)
        section.bottom_margin = Cm(cm)
        section.left_margin = Cm(cm)
        section.right_margin = Cm(cm)


def _add_centered(doc: Document, text: str, bold: bool = False, size_pt: float = 12.0) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(size_pt)


def _add_body(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)


def _make_correct_doc(path: Path) -> None:
    """Tüm kurallara uygun örnek resmi yazı oluşturur."""
    doc = Document()
    _set_margins(doc, cm=1.5)

    # Başlık bloğu
    _add_centered(doc, "T.C.", bold=True)
    _add_centered(doc, "GEBZE TEKNİK ÜNİVERSİTESİ REKTÖRLÜĞÜ", bold=True)
    _add_centered(doc, "MÜHENDİSLİK FAKÜLTESİ", bold=True)

    # Sayı / Tarih
    p = doc.add_paragraph()
    run = p.add_run("Sayı: GTU-2025-001                                           Tarih: 26.05.2025")
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)

    # Konu
    p = doc.add_paragraph()
    run = p.add_run("Konu: Test Yazısı")
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)

    # Muhatap
    p = doc.add_paragraph()
    run = p.add_run("Rektörlük Makamına,")
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)

    # Metin gövdesi
    _add_body(
        doc,
        "Fakültemiz bünyesinde yürütülen akademik faaliyetlere ilişkin bilgiler "
        "ekte sunulmaktadır. Gereğini bilgilerinize arz ederim.",
    )

    # Kapanış
    p = doc.add_paragraph()
    run = p.add_run("Arz ederim.")
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)

    # İmza
    p = doc.add_paragraph()
    run = p.add_run("Prof. Dr. Ahmet Yılmaz\nDekan")
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)

    doc.save(str(path))


def _make_errors_doc(path: Path) -> None:
    """Kasıtlı format ve alan hataları içeren belge oluşturur."""
    doc = Document()
    # Yanlış marjlar: 3 cm
    _set_margins(doc, cm=3.0)

    # T.C. YOK — FLD-001 tetiklenir
    # Üniversite adı yok
    # Birim adı yok

    # Sayı: tarih yok → FLD-004 tetiklenir
    p = doc.add_paragraph()
    run = p.add_run("Sayı: GTU-2025-999")
    run.font.name = "Calibri"
    run.font.size = Pt(11)

    # Konu: var — parser'ın header_done veya konu_done yapması için gerekli
    p = doc.add_paragraph()
    run = p.add_run("Konu: Hatalı Test Yazısı")
    run.font.name = "Calibri"
    run.font.size = Pt(11)

    # Metin — Calibri kullanılmış → FMT-001; çift boşluk ve virgül öncesi boşluk → LNG kuralları
    p = doc.add_paragraph()
    run = p.add_run("Bu  belgede  çift  boşluk ve virgül öncesi boşluk , var.")
    run.font.name = "Calibri"
    run.font.size = Pt(11)

    # Kapanış YOK → CLS-001 tetiklenir
    # İmza da yok

    doc.save(str(path))


# ── Fikstürler ────────────────────────────────────────────────────────────────

@pytest.fixture(scope="session")
def correct_docx_path(tmp_path_factory) -> Path:
    path = tmp_path_factory.mktemp("docs") / "test_correct.docx"
    _make_correct_doc(path)
    return path


@pytest.fixture(scope="session")
def errors_docx_path(tmp_path_factory) -> Path:
    path = tmp_path_factory.mktemp("docs") / "test_errors.docx"
    _make_errors_doc(path)
    return path


@pytest.fixture(scope="session")
def rag_retriever(tmp_path_factory):
    """RAG retriever fikstürü — kaynak veri yoksa testi atlar."""
    try:
        from app.rag.retriever import GuidelineRetriever
        from app.rag.ingest import run_ingest
    except ImportError:
        pytest.skip("RAG modülleri bulunamadı")

    data_dir = Path(__file__).resolve().parents[2] / "data"
    tdk_json = data_dir / "tdk_rules.json"
    if not tdk_json.exists():
        pytest.skip("data/tdk_rules.json bulunamadı — RAG testleri atlanıyor")

    chroma_dir = tmp_path_factory.mktemp("chroma")
    run_ingest(data_dir=str(data_dir), persist_dir=str(chroma_dir))
    retriever = GuidelineRetriever(persist_dir=str(chroma_dir))
    if not retriever.is_ready:
        pytest.skip("RAG retriever hazır değil")
    return retriever
